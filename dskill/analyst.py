"""Trace-contrast analyst, step 1 of the D-Skill pipeline.

One LLM call per task compares an expert and a student trace of the same task
(no answer key, no judge) and writes Failure/Success Memory Items, parsed into
the upstream formats (parsed_errors.json, parsed_success.json).

    python -m dskill.analyst --prompt <name> \
        --student_logs ... --expert_logs ... --output_dir ... \
        --instance_ids ... --model ... --parsed_errors ... --parsed_success ...

--prompt is a key of _PROMPTS; the alfworld and chess prompts also take
--items_json.
"""

import argparse
import json
import re
import sys
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import dskill  # noqa: E402,F401  (puts third_party/trace2skill on sys.path)

from react_agent.models import Message  # noqa: E402


def _rep(s, old, new):
    assert old in s, old[:60]
    return s.replace(old, new)


# spreadsheet: expert and student traces, each followed by its answer range.
PROMPT_SPREADSHEET = """\
You are an expert coach for spreadsheet-manipulation agents. You will \
see two execution traces for the SAME task, produced under the same \
ReAct protocol (a bash tool, alternating Action/Observation turns, \
ending with ACTION: TASK_COMPLETE).

- The EXPERT trace comes from a much stronger agent whose working style \
the student should learn to imitate.
- The STUDENT trace comes from the weaker agent you are coaching.

You have NO answer key and NO correctness information for either trace, \
and you must not invent any. Do not judge whether either agent solved \
the task. Your job is purely behavioral: identify how the expert's way \
of working differs from the student's, and convert the most \
consequential differences into reusable lessons that would make the \
student's future traces resemble the expert's.

FIRST, before writing any items: diff the two episodes' FINAL artifacts \
-- state concretely what the expert's output contains that the \
student's does not (or vice versa), and trace each output difference \
back to the decision that caused it. At least one item must come from \
this output delta rather than from workflow style.

Compare along dimensions such as: how each agent INTERPRETED the \
instruction -- what deliverable the grader expects, and how ambiguous \
phrasing was construed (e.g., 'create a macro/formula that does X' \
usually grades the RESULT of X, not the code); what each agent \
inspected before acting; how each chose its approach; formulas vs \
computed values in written cells; how the output workbook was produced \
at the exact required path; verification steps before declaring \
completion; recovery behavior after a failed or empty observation; \
number and economy of turns; completion discipline.

Each trace is followed by the ANSWER-RANGE CONTENTS of the workbook \
that agent produced — the exact cells the grader reads. Where the two \
differ, trace the difference back to the DECISION or INSPECTION STEP \
that produced it and write the general habit that would have avoided \
it. NEVER state what the correct value for this task was, and never \
write a rule about this task's particular columns, sheets, or values: \
the student you are coaching will face different tasks, and a rule \
keyed to this one is worthless to it. You still have NO correctness \
information — a difference does not tell you which agent is right.

Anchor every observation in specific actions visible in the traces — \
quote or paraphrase the concrete step. Never speculate about steps that \
are not shown.

Write AT MOST 3 items total, ordered by expected impact on the \
student's future behavior:
- "Failure Memory Item" — a behavioral gap: what the student did, what \
the expert did instead, and ONE concrete reusable instruction.
- at most 1 "Success Memory Item" — a behavior the student already \
shares with the expert and should keep.

Every item's Content must additionally include one concrete check the \
student can EXECUTE at runtime without any answer key (e.g., 'after \
writing, confirm the answer range holds data values, not code/formula \
text, unless the task explicitly grades formulas').

Every instruction must generalize across tasks: procedures and habits \
("always inspect every sheet's headers before writing"), never \
task-specific values, sheet names, or cell addresses.

Output format (exactly this structure, failure items first):

# Failure Memory Item 1
## Title
<short name>
## Description
<the observed student-vs-expert difference, with trace evidence>
## Content
<the single reusable instruction the student should follow>

# Success Memory Item 1
## Title
<short name>
## Description
<the shared behavior, with trace evidence>
## Content
<the instruction to keep doing it>
"""

# chess base: expert and student solving traces under the typed-tool protocol;
# chess_commit_check and chess_fallible_verify derive from it.
PROMPT_CHESS = """\
You are an expert coach for a chess tactics-solving agent that uses \
typed analysis tools. You will see two solving traces for the SAME \
puzzle. Each turn, an agent sees the position (FEN, printed board, \
legal-move list) and may either call ONE analysis tool -- <hanging>, \
<threats>, <safety>, <simulate> (one candidate vs up to 5 replies in \
one call), <board> (position after one own move), <line> (replay of a \
full calculated line, own moves and opponent replies alternating) -- \
and read its result; the material count is printed free in every \
observation, or commit a move in <action>...</action> tags; a per-move \
budget limits tool calls.

- The EXPERT trace comes from a much stronger player whose working \
style the student should learn to imitate.
- The STUDENT trace comes from the weaker player you are coaching.

You have NO answer key and NO correctness information for either trace, \
and you must not invent any. Do not judge whether either agent solved \
the puzzle, and do NOT assert that any specific move was the correct \
answer. Your job is purely behavioral: identify how the expert's way of \
working differs from the student's, and convert the most consequential \
differences into reusable lessons that would make the student's future \
traces resemble the expert's.

Compare along dimensions such as: how each player READ THE POSITION \
before generating candidates -- what kind of position each treated it \
as (tactical melee, quiet/positional maneuvering, or an endgame), which \
structural features each named (pawn structure and majorities, king \
safety, piece activity), and whether each registered threats against \
ITS OWN side (own pieces attacked or overloaded) before attacking; how \
each player GENERATED candidate moves (enumerating checks, captures, \
and threats from the legal-move list; naming the tactical motif -- \
fork, pin, skewer, discovered attack, back rank, overloaded defender -- \
that pointed to a candidate); which tools each called, in what order, \
and what each actually did with the results; CALIBRATION in reading \
tool output -- e.g. discarding a candidate because of a surface warning \
(an exchange-risk flag) without simulating the follow-up, trusting a \
mental picture over a tool result, or banking a small verified gain \
without checking a stronger forcing candidate; how each compared \
multiple promising candidates before committing; how each converted the \
second move of a line after the opponent's reply; economy of the tool \
budget.

Anchor every observation in specific turns visible in the traces -- \
quote or paraphrase the concrete tool call or move. Never speculate \
about turns that are not shown.

Write AT MOST 3 items total, ordered by expected impact on the \
student's future behavior:
- "Failure Memory Item" -- a behavioral gap: what the student did, what \
the expert did instead, and ONE concrete reusable instruction.
- at most 1 "Success Memory Item" -- a behavior the student already \
shares with the expert and should keep.

Every item's Content must additionally include one concrete check the \
student can perform at runtime without any answer key.

Every instruction must generalize across puzzles: procedures, habits, \
and ways of reading tool output -- never this puzzle's specific \
squares, piece placements, or move names.

Output format (exactly this structure, failure items first):

# Failure Memory Item 1
## Title
<short name>
## Description
<the observed student-vs-expert difference, with trace evidence>
## Content
<the single reusable instruction the student should follow>

# Success Memory Item 1
## Title
<short name>
## Description
<the shared behavior, with trace evidence>
## Content
<the instruction to keep doing it>
"""

# chess_commit_check: PROMPT_CHESS plus commitment and abandonment dimensions.
# for Haiku -> Qwen, the experiments are with text-tag harnesss, so the prompt is written to encourage commit and abandon checks
PROMPT_CHESS_COMMIT_CHECK = _rep(
    PROMPT_CHESS,
    "after the opponent's reply; ",
    "after the opponent's reply; "
    "how each player's FINAL COMMITMENT was backed -- what tool result in "
    "the current position supported the move ultimately committed (or the "
    "final turn, if none was), and whether concluding that a move was best "
    "was immediately followed by exactly one action tag in the final "
    "channel; whether either player ABANDONED a candidate it had itself "
    "priced as winning or verified as forcing without a refuting tool "
    "result from the current position -- what the pricing turn and the "
    "abandonment turn each showed, and what replaced the candidate; ")

# chess_fallible_verify: chess plus a verification-before-commit dimension,
# with the expert framed as fallible and moves judged on board evidence.
# for Sonnet -> Haiku with the claude code harness
PROMPT_CHESS_FALLIBLE_VERIFY = _rep(
    PROMPT_CHESS,
    "before committing; how each converted the second move",
    "before committing; VERIFICATION before committing -- whether each "
    "player followed its chosen line to a concrete end (checkmate, or the "
    "material count after the opponent's forced reply) using its tools "
    "before committing the move; a commitment made without that "
    "verification is a gap to report even when the expert made it, and "
    "never a behavior to imitate; how each converted the second move")
PROMPT_CHESS_FALLIBLE_VERIFY = _rep(
    PROMPT_CHESS_FALLIBLE_VERIFY,
    """- The EXPERT trace comes from a much stronger player whose working \
style the student should learn to imitate.
- The STUDENT trace comes from the weaker player you are coaching.
""",
    """- The EXPERT trace comes from a stronger player. Its working style is \
usually worth imitating, but its moves are not always right.
- The STUDENT trace comes from the weaker player you are coaching.

Judge each move by the evidence in the traces (checkmate, material after \
the forced reply, tool results), not by who played it. When the student's \
move is better supported by that evidence, keep it as a Success item; \
never turn an expert move the evidence contradicts into a lesson.
""")

# alfworld: the same coaching frame for the ALFWorld step protocol.
PROMPT_ALFWORLD = """\
You are an expert coach for embodied household-task agents operating \
in the ALFWorld text environment. You will see two episode traces for \
the SAME task. Each step, an agent sees the current observation and a \
list of admissible actions, and must emit one action in \
<action>...</action> tags; episodes end on success or after the step \
budget.

- The EXPERT trace comes from a much stronger agent whose working \
style the student should learn to imitate.
- The STUDENT trace comes from the weaker agent you are coaching.

You have NO answer key and NO correctness information for either \
trace, and you must not invent any. Do not judge whether either agent \
solved the task. Your job is purely behavioral: identify how the \
expert's way of working differs from the student's, and convert the \
most consequential differences into reusable lessons that would make \
the student's future traces resemble the expert's.

Compare along dimensions such as: how each agent construed the goal \
and its implied preparation steps; what each inspected before acting; \
how each searched for the target object; whether actions were chosen \
from the admissible list; how each reacted to a failed or \
uninformative observation; number and economy of steps; completion \
discipline.

Anchor every observation in specific actions visible in the traces -- \
quote or paraphrase the concrete step. Never speculate about steps \
that are not shown.

Write AT MOST 3 items total, ordered by expected impact on the \
student's future behavior:
- "Failure Memory Item" -- a behavioral gap: what the student did, \
what the expert did instead, and ONE concrete reusable instruction.
- at most 1 "Success Memory Item" -- a behavior the student already \
shares with the expert and should keep.

Every item's Content must additionally include one concrete check the \
student can perform at runtime without any answer key.

Every instruction must generalize across episodes: procedures and \
habits, never this episode's specific room, object instances, or \
receptacle numbers.

Output format (exactly this structure, failure items first):

# Failure Memory Item 1
## Title
<short name>
## Description
<the observed student-vs-expert difference, with trace evidence>
## Content
<the single reusable instruction the student should follow>

# Success Memory Item 1
## Title
<short name>
## Description
<the shared behavior, with trace evidence>
## Content
<the instruction to keep doing it>
"""
assert "PRIORITY RUBRIC" not in PROMPT_ALFWORLD
assert "STANDING" not in PROMPT_ALFWORLD
assert "## Criticality" not in PROMPT_ALFWORLD
assert "ALFWorld" in PROMPT_ALFWORLD

# spreadsheet_teacher_only: PROMPT_SPREADSHEET rewritten for the expert alone.
PROMPT_SPREADSHEET_TEACHER_ONLY = PROMPT_SPREADSHEET
PROMPT_SPREADSHEET_TEACHER_ONLY = _rep(
    PROMPT_SPREADSHEET_TEACHER_ONLY,
    "You will see two execution traces for the SAME task, produced under "
    "the same ReAct protocol",
    "You will see ONE execution trace of an EXPERT agent, produced under a "
    "ReAct protocol")
PROMPT_SPREADSHEET_TEACHER_ONLY = _rep(
    PROMPT_SPREADSHEET_TEACHER_ONLY,
    "- The EXPERT trace comes from a much stronger agent whose working "
    "style the student should learn to imitate.\n"
    "- The STUDENT trace comes from the weaker agent you are coaching.",
    "The trace comes from a much stronger agent whose working style a "
    "weaker student should learn to imitate. You never see the student; "
    "distill the expert's most consequential reusable behaviors into "
    "lessons the student can follow.")
PROMPT_SPREADSHEET_TEACHER_ONLY = _rep(
    PROMPT_SPREADSHEET_TEACHER_ONLY,
    "Do not judge whether either agent solved the task. Your job is purely "
    "behavioral: identify how the expert's way of working differs from the "
    "student's, and convert the most consequential differences into "
    "reusable lessons that would make the student's future traces resemble "
    "the expert's.",
    "Do not judge whether the expert solved the task. Your job is purely "
    "behavioral: identify the expert's consequential working habits and "
    "convert them into reusable lessons that would make a weaker student's "
    "future traces resemble the expert's.")
PROMPT_SPREADSHEET_TEACHER_ONLY = _rep(
    PROMPT_SPREADSHEET_TEACHER_ONLY,
    "diff the two episodes' FINAL artifacts -- state concretely what the "
    "expert's output contains that the student's does not (or vice versa), "
    "and trace each output difference back to the decision that caused it. "
    "At least one item must come from this output delta rather than from "
    "workflow style.",
    "examine the expert's FINAL artifact -- state concretely what its "
    "output contains, and trace each notable property back to the decision "
    "or inspection step that produced it. At least one item must come from "
    "this artifact rather than from workflow style.")
PROMPT_SPREADSHEET_TEACHER_ONLY = _rep(
    PROMPT_SPREADSHEET_TEACHER_ONLY,
    "how each agent INTERPRETED the instruction -- what deliverable the "
    "grader expects, and how ambiguous phrasing was construed (e.g., "
    "'create a macro/formula that does X' usually grades the RESULT of X, "
    "not the code); what each agent inspected before acting; how each "
    "chose its approach;",
    "how the expert INTERPRETED the instruction -- what deliverable the "
    "grader expects, and how ambiguous phrasing was construed (e.g., "
    "'create a macro/formula that does X' usually grades the RESULT of X, "
    "not the code); what the expert inspected before acting; how it chose "
    "its approach;")
PROMPT_SPREADSHEET_TEACHER_ONLY = _rep(
    PROMPT_SPREADSHEET_TEACHER_ONLY,
    "Each trace is followed by the ANSWER-RANGE CONTENTS of the workbook "
    "that agent produced — the exact cells the grader reads. Where the two "
    "differ, trace the difference back to the DECISION or INSPECTION STEP "
    "that produced it and write the general habit that would have avoided "
    "it. NEVER state what the correct value for this task was, and never "
    "write a rule about this task's particular columns, sheets, or values: "
    "the student you are coaching will face different tasks, and a rule "
    "keyed to this one is worthless to it. You still have NO correctness "
    "information — a difference does not tell you which agent is right.",
    "The trace is followed by the ANSWER-RANGE CONTENTS of the workbook "
    "the expert produced — the exact cells the grader reads. Trace each "
    "notable property of that output back to the DECISION or INSPECTION "
    "STEP that produced it and write the general habit behind it. NEVER "
    "state what the correct value for this task was, and never write a "
    "rule about this task's particular columns, sheets, or values: the "
    "student you are coaching will face different tasks, and a rule keyed "
    "to this one is worthless to it. You still have NO correctness "
    "information about this output.")
PROMPT_SPREADSHEET_TEACHER_ONLY = _rep(
    PROMPT_SPREADSHEET_TEACHER_ONLY,
    "\"Failure Memory Item\" — a behavioral gap: what the student did, "
    "what the expert did instead, and ONE concrete reusable instruction.\n"
    "- at most 1 \"Success Memory Item\" — a behavior the student already "
    "shares with the expert and should keep.",
    "\"Failure Memory Item\" — an expert habit a weaker student would "
    "plausibly lack, and ONE concrete reusable instruction.\n"
    "- at most 1 \"Success Memory Item\" — a foundational habit so basic "
    "the student likely has it already and must keep it.")
PROMPT_SPREADSHEET_TEACHER_ONLY = _rep(
    PROMPT_SPREADSHEET_TEACHER_ONLY,
    "<the observed student-vs-expert difference, with trace evidence>",
    "<the observed expert behavior, with trace evidence>")
PROMPT_SPREADSHEET_TEACHER_ONLY = _rep(
    PROMPT_SPREADSHEET_TEACHER_ONLY,
    "<the shared behavior, with trace evidence>",
    "<the habit, with trace evidence>")
assert "STUDENT trace" not in PROMPT_SPREADSHEET_TEACHER_ONLY

# spreadsheet_student_only: the same rewrite for the student alone.
PROMPT_SPREADSHEET_STUDENT_ONLY = PROMPT_SPREADSHEET
PROMPT_SPREADSHEET_STUDENT_ONLY = _rep(
    PROMPT_SPREADSHEET_STUDENT_ONLY,
    "You will see two execution traces for the SAME task, produced under "
    "the same ReAct protocol",
    "You will see ONE execution trace of a STUDENT agent, produced under a "
    "ReAct protocol")
PROMPT_SPREADSHEET_STUDENT_ONLY = _rep(
    PROMPT_SPREADSHEET_STUDENT_ONLY,
    "- The EXPERT trace comes from a much stronger agent whose working "
    "style the student should learn to imitate.\n"
    "- The STUDENT trace comes from the weaker agent you are coaching.",
    "The trace comes from the weaker agent you are coaching. You never see "
    "an expert; identify the student's consequential working habits -- the "
    "ones most likely to decide whether its output is right -- and distill "
    "them into lessons that would make its future traces more reliable.")
PROMPT_SPREADSHEET_STUDENT_ONLY = _rep(
    PROMPT_SPREADSHEET_STUDENT_ONLY,
    "Do not judge whether either agent solved the task. Your job is purely "
    "behavioral: identify how the expert's way of working differs from the "
    "student's, and convert the most consequential differences into "
    "reusable lessons that would make the student's future traces resemble "
    "the expert's.",
    "Do not judge whether the student solved the task. Your job is purely "
    "behavioral: identify the student's consequential working habits, flag "
    "the ones that look fragile (unchecked assumptions, skipped "
    "inspections, unverified writes), and convert them into reusable "
    "lessons that would make the student's future traces more careful and "
    "more robust.")
PROMPT_SPREADSHEET_STUDENT_ONLY = _rep(
    PROMPT_SPREADSHEET_STUDENT_ONLY,
    "diff the two episodes' FINAL artifacts -- state concretely what the "
    "expert's output contains that the student's does not (or vice versa), "
    "and trace each output difference back to the decision that caused it. "
    "At least one item must come from this output delta rather than from "
    "workflow style.",
    "examine the student's FINAL artifact -- state concretely what its "
    "output contains, and trace each notable property back to the decision "
    "or inspection step that produced it. At least one item must come from "
    "this artifact rather than from workflow style.")
PROMPT_SPREADSHEET_STUDENT_ONLY = _rep(
    PROMPT_SPREADSHEET_STUDENT_ONLY,
    "how each agent INTERPRETED the instruction -- what deliverable the "
    "grader expects, and how ambiguous phrasing was construed (e.g., "
    "'create a macro/formula that does X' usually grades the RESULT of X, "
    "not the code); what each agent inspected before acting; how each "
    "chose its approach;",
    "how the student INTERPRETED the instruction -- what deliverable the "
    "grader expects, and how ambiguous phrasing was construed (e.g., "
    "'create a macro/formula that does X' usually grades the RESULT of X, "
    "not the code); what the student inspected before acting; how it chose "
    "its approach;")
PROMPT_SPREADSHEET_STUDENT_ONLY = _rep(
    PROMPT_SPREADSHEET_STUDENT_ONLY,
    "Each trace is followed by the ANSWER-RANGE CONTENTS of the workbook "
    "that agent produced — the exact cells the grader reads. Where the two "
    "differ, trace the difference back to the DECISION or INSPECTION STEP "
    "that produced it and write the general habit that would have avoided "
    "it. NEVER state what the correct value for this task was, and never "
    "write a rule about this task's particular columns, sheets, or values: "
    "the student you are coaching will face different tasks, and a rule "
    "keyed to this one is worthless to it. You still have NO correctness "
    "information — a difference does not tell you which agent is right.",
    "The trace is followed by the ANSWER-RANGE CONTENTS of the workbook "
    "the student produced — the exact cells the grader reads. Trace each "
    "notable property of that output back to the DECISION or INSPECTION "
    "STEP that produced it and write the general habit behind it; where a "
    "property looks like the product of a shortcut (an unchecked "
    "assumption, a skipped inspection, a value typed instead of computed), "
    "say so and write the corrective habit. NEVER state what the correct "
    "value for this task was, and never write a rule about this task's "
    "particular columns, sheets, or values: the student will face "
    "different tasks, and a rule keyed to this one is worthless to it. You "
    "still have NO correctness information about this output.")
PROMPT_SPREADSHEET_STUDENT_ONLY = _rep(
    PROMPT_SPREADSHEET_STUDENT_ONLY,
    "\"Failure Memory Item\" — a behavioral gap: what the student did, "
    "what the expert did instead, and ONE concrete reusable instruction.\n"
    "- at most 1 \"Success Memory Item\" — a behavior the student already "
    "shares with the expert and should keep.",
    "\"Failure Memory Item\" — a fragile habit in the student's way of "
    "working that plausibly costs it tasks, and ONE concrete reusable "
    "instruction.\n"
    "- at most 1 \"Success Memory Item\" — a sound habit the student "
    "already has and must keep.")
PROMPT_SPREADSHEET_STUDENT_ONLY = _rep(
    PROMPT_SPREADSHEET_STUDENT_ONLY,
    "<the observed student-vs-expert difference, with trace evidence>",
    "<the observed student behavior and why it is fragile, with trace "
    "evidence>")
PROMPT_SPREADSHEET_STUDENT_ONLY = _rep(
    PROMPT_SPREADSHEET_STUDENT_ONLY,
    "<the shared behavior, with trace evidence>",
    "<the habit, with trace evidence>")
assert "EXPERT trace" not in PROMPT_SPREADSHEET_STUDENT_ONLY


_PROMPTS = {
    "spreadsheet": PROMPT_SPREADSHEET,
    "spreadsheet_teacher_only": PROMPT_SPREADSHEET_TEACHER_ONLY,
    "spreadsheet_student_only": PROMPT_SPREADSHEET_STUDENT_ONLY,
    "alfworld": PROMPT_ALFWORLD,
    "chess_commit_check": PROMPT_CHESS_COMMIT_CHECK,
    "chess_fallible_verify": PROMPT_CHESS_FALLIBLE_VERIFY,
}


def _select_system(name: str) -> str:
    if name not in _PROMPTS:
        raise SystemExit(
            f"unknown analyst prompt {name!r}; valid: " + ", ".join(_PROMPTS))
    return _PROMPTS[name]


def _answer_ranges(ap: str):
    """Parse an answer_position into [(sheet, c1, r1, c2, r2), ...]."""
    out = []
    for part in str(ap or "").split(","):
        m = re.match(r"\s*(?:'?([^'!]+)'?!)?\$?([A-Z]+)\$?(\d+)"
                     r"(?::\$?([A-Z]+)\$?(\d+))?\s*$", part)
        if m:
            sh, c1, r1, c2, r2 = m.groups()
            out.append((sh, c1, int(r1), c2 or c1, int(r2 or r1)))
    return out


def _answer_range_dump(run_dir: Path, tid: str, rec: dict,
                       limit: int = 4000, head: int = 40, tail: int = 10) -> str:
    """Dump the produced workbook's answer-range cells (capped, elided)."""
    import warnings
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        try:
            import openpyxl
            from openpyxl.utils import column_index_from_string as ci
        except Exception:                                   # noqa: BLE001
            return "(openpyxl unavailable)"
        xs = sorted((run_dir / "outputs" / "spreadsheet" / tid).glob("*_output.xlsx")) \
            if (run_dir / "outputs" / "spreadsheet" / tid).exists() else []
        if not xs:
            return "(no output workbook was produced)"
        try:
            wb = openpyxl.load_workbook(xs[0], data_only=True)
        except Exception as e:                              # noqa: BLE001
            return f"(could not read output workbook: {e})"
        rngs = _answer_ranges(rec.get("answer_position"))
        lines, elided = [], 0
        if not rngs:                                        # fallback: used range
            ws = wb.active
            rngs = [(ws.title, "A", 1, ws.cell(1, ws.max_column).column_letter,
                     ws.max_row)]
        for sh, c1, r1, c2, r2 in rngs:
            name = sh or rec.get("answer_sheet")
            ws = wb[name] if name in wb.sheetnames else wb.active
            lines.append(f"--- {ws.title}!{c1}{r1}:{c2}{r2}")
            rows = list(range(r1, min(r2, ws.max_row) + 1))
            keep = rows if len(rows) <= head + tail else rows[:head] + rows[-tail:]
            cut_at = head if len(rows) > head + tail else -1
            for i, r in enumerate(keep):
                if i == cut_at:
                    elided = len(rows) - head - tail
                    lines.append(f"    ... ({elided} rows elided) ...")
                vals = tuple(ws.cell(r, c).value
                             for c in range(ci(c1), ci(c2) + 1))
                lines.append(f"{r}: {vals!r}")
        text = "\n".join(lines)
        if len(text) > limit:
            text = text[:limit] + f"\n    ... (truncated at {limit} chars) ..."
        return text


def _run_dir(logs_dir: Path, tid: str) -> Path:
    """Run dir of tid's produced workbook (resolves merged-log symlinks)."""
    p = next(iter(logs_dir.glob(f"*_{tid}.md")), None)
    return p.resolve().parent.parent if p else logs_dir.parent


def _trace_tail(path: Path, limit: int = 35000,
                head: int = 10000) -> str:
    """Trace from the task turn on, middle elided beyond head+limit chars."""
    if path is None or not path.exists():
        return "(trace missing)"
    text = path.read_text()
    # drop the system-prompt preamble; keep from the task turn onward
    m = re.search(r"^## \[2\] USER", text, re.MULTILINE)
    if m:
        text = text[m.start():]
    if len(text) <= limit + head:
        return text
    elided = len(text) - head - limit
    return (text[:head] + f"\n    ... ({elided} chars elided) ...\n"
            + text[-limit:])


def _contrast_one(tid: str, rec: dict, student_logs: Path,
                  expert_logs: Path, out_dir: Path, client_factory,
                  prompt: str):
    cdir = out_dir / tid
    cdir.mkdir(parents=True, exist_ok=True)
    report_path = cdir / "contrast_report.md"
    if report_path.exists():
        return                                   # resume: already analyzed
    stu = _trace_tail(next(iter(student_logs.glob(f"*_{tid}.md")), None))
    exp = _trace_tail(next(iter(expert_logs.glob(f"*_{tid}.md")), None))
    system = _select_system(prompt)
    if prompt == "spreadsheet":
        ap = rec.get("answer_position")
        exp_wb = _answer_range_dump(_run_dir(expert_logs, tid), tid, rec)
        stu_wb = _answer_range_dump(_run_dir(student_logs, tid), tid, rec)
        user = (f"### task instruction\n{rec['instruction']}\n\n"
                f"### graded answer range\n{ap}\n\n"
                f"### EXPERT trace (behavioral target)\n{exp}\n\n"
                f"### EXPERT produced answer range\n{exp_wb}\n\n"
                f"### STUDENT trace\n{stu}\n\n"
                f"### STUDENT produced answer range\n{stu_wb}\n\n"
                "Compare the two traces and the two produced answer ranges, "
                "then write your items now, in the exact format from your "
                "instructions.")
    elif prompt == "spreadsheet_teacher_only":
        ap = rec.get("answer_position")
        exp_wb = _answer_range_dump(_run_dir(expert_logs, tid), tid, rec)
        user = (f"### task instruction\n{rec['instruction']}\n\n"
                f"### graded answer range\n{ap}\n\n"
                f"### EXPERT trace (behavioral target)\n{exp}\n\n"
                f"### EXPERT produced answer range\n{exp_wb}\n\n"
                "Distill the expert's reusable behaviors and write your "
                "items now, in the exact format from your instructions.")
    elif prompt == "spreadsheet_student_only":
        ap = rec.get("answer_position")
        stu_wb = _answer_range_dump(_run_dir(student_logs, tid), tid, rec)
        user = (f"### task instruction\n{rec['instruction']}\n\n"
                f"### graded answer range\n{ap}\n\n"
                f"### STUDENT trace\n{stu}\n\n"
                f"### STUDENT produced answer range\n{stu_wb}\n\n"
                "Examine the student's trace and its produced answer range, "
                "then write your items now, in the exact format from your "
                "instructions.")
    else:
        user = (f"### task instruction\n{rec['instruction']}\n\n"
                f"### EXPERT trace (behavioral target)\n{exp}\n\n"
                f"### STUDENT trace\n{stu}\n\n"
                "Compare the two traces and write your items now, in the "
                "exact format from your instructions.")
    (cdir / "analyst_prompt.md").write_text(
        f"# SYSTEM\n{system}\n\n# USER\n{user}")
    try:
        reply = client_factory().chat([
            Message(role="system", content=system),
            Message(role="user", content=user)])
    except Exception as e:                        # noqa: BLE001
        print(f"[contrast] {tid} error: {e}", file=sys.stderr)
        reply = ""
    report_path.write_text(reply or "(empty analyst reply)")
    print(f"[contrast] {tid} done", flush=True)


def main():
    from dskill.llm.pin_provider import install as _pin_install
    _pin_install()
    ap = argparse.ArgumentParser()
    ap.add_argument("--prompt", required=True, choices=list(_PROMPTS),
                    help="spreadsheet (SpreadsheetBench), "
                         "spreadsheet_teacher_only / spreadsheet_student_only "
                         "(one-sided traces), alfworld (ALFWorld), "
                         "chess_fallible_verify (Sonnet -> Haiku) or "
                         "chess_commit_check (Haiku -> Qwen)")
    ap.add_argument("--student_logs", required=True)
    ap.add_argument("--expert_logs", required=True)
    ap.add_argument("--output_dir", required=True)
    ap.add_argument("--instance_ids", required=True)
    ap.add_argument("--items_json", default=None,
                    help="items.json for non-spreadsheet domains "
                         "(ids ':'->'-' normalized)")
    ap.add_argument("--model", required=True)
    ap.add_argument("--parsed_errors", required=True)
    ap.add_argument("--parsed_success", required=True)
    ap.add_argument("--workers", type=int, default=8)
    args = ap.parse_args()

    from dskill.llm.client import make_client
    if getattr(args, "items_json", None):
        _items = json.loads(Path(args.items_json).expanduser().read_text())
        by_id = {}
        for _it in _items:
            _iid = str(_it.get("uid") or _it["id"]).replace(":", "-")
            _ch = "\n".join(f"{c['label']}. {c['text']}"
                             for c in _it.get("choices", []))
            by_id[_iid] = {"id": _iid,
                           "instruction": (_it.get("question", "")
                                           + "\n\nChoices:\n" + _ch)}
    else:
        from envs.spreadsheetbench.dataset import load_verified
        _, by_id = load_verified()
    ids = [t for t in args.instance_ids.split(",") if t.strip()]
    out_dir = Path(args.output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    student_logs, expert_logs = Path(args.student_logs), Path(args.expert_logs)

    with ThreadPoolExecutor(max_workers=args.workers) as ex:
        list(ex.map(lambda t: _contrast_one(
            t, by_id[t], student_logs, expert_logs, out_dir,
            lambda: make_client(args.model), args.prompt), ids))

    # Split at the Success item; the upstream failure parser would swallow it.
    from analysis.parse_error_analysis_outputs import parse_items as parse_fail
    from analysis.parse_success_analysis_outputs import parse_items as parse_succ
    errors, successes = [], []
    for tid in ids:
        rp = out_dir / tid / "contrast_report.md"
        if not rp.exists():
            continue
        text = rp.read_text()
        m = re.search(r"^#\s+Success Memory Item", text, re.MULTILINE)
        fail_half = text[:m.start()] if m else text
        succ_half = text[m.start():] if m else ""
        fitems = parse_fail(fail_half)
        sitems = parse_succ(succ_half) if succ_half else []
        if fitems:
            errors.append({"instance_id": tid,
                           "source_file": "contrast_report.md",
                           "items": fitems})
        if sitems:
            successes.append({"instance_id": tid,
                              "source_file": "contrast_report.md",
                              "items": sitems})
    Path(args.parsed_errors).write_text(json.dumps(errors, indent=1))
    Path(args.parsed_success).write_text(json.dumps(successes, indent=1))
    n_f = sum(len(r["items"]) for r in errors)
    n_s = sum(len(r["items"]) for r in successes)
    print(f"[contrast] parsed: {n_f} gap items over {len(errors)} tasks, "
          f"{n_s} keep items over {len(successes)} tasks", flush=True)


if __name__ == "__main__":
    main()
